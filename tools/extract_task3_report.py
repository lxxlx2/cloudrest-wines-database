#!/usr/bin/env python3
"""Extract the current final report from Task 3 through its last appendix."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import sys

ROOT=Path(__file__).resolve().parents[1]
src=ROOT/"deliverables/final-submission/Cloudrest_Wines_Report.docx"
out=ROOT/"deliverables/final-submission/Cloudrest_Wines_Report_Task3_Onwards.docx"
doc=Document(src)
body=doc.element.body
start=None
for i,element in enumerate(body):
    if element.tag!=qn("w:p"):continue
    t="".join(n.text or "" for n in element.iter(qn("w:t"))).strip()
    ppr=element.find(qn("w:pPr"))
    st=ppr.find(qn("w:pStyle")) if ppr is not None else None
    style=st.get(qn("w:val"),"") if st is not None else ""
    if t.startswith(("Task 3 —","Task 3 -","Task 3:")) and ("Heading" in style or not style):
        start=i;break
if start is None:sys.exit("Task 3 heading not found in final Word")
for item in list(body)[:start]:body.remove(item)
doc.save(out)
check=Document(out)
assert check.paragraphs[0].text.startswith("Task 3"),check.paragraphs[0].text
assert any(x.text.startswith("Task 7") for x in check.paragraphs)
assert len(check.tables)>20
print(f"EXTRACTED={out} PARAGRAPHS={len(check.paragraphs)} TABLES={len(check.tables)}")
